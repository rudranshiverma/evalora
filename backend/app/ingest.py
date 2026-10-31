import io
import re
from dataclasses import dataclass
from pathlib import PurePosixPath
import docx
import pymupdf

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
SHORT_SUBMISSION_WORDS = 30

class ExtractionError(Exception):
    pass

@dataclass(frozen=True)
class Extracted:
    text: str
    word_count: int
    warning: str | None = None

def _path(filename: str) -> PurePosixPath:
    return PurePosixPath(filename.replace("\\", "/"))

def extract_text(filename: str, data: bytes) -> Extracted:
    suffix = _path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        kind = suffix or "with no extension"
        return Extracted("", 0, f"Unsupported file type {kind}. Use PDF, DOCX or TXT.")
    try:
        if suffix == ".pdf":
            raw = _pdf_text(data)
        elif suffix == ".docx":
            raw = _docx_text(data)
        else:
            raw = _decode(data)
    except ExtractionError as e:
        return Extracted("", 0, str(e))
    except Exception:
        return Extracted("", 0, "Could not read this file. It may be corrupted.")

    text = clean_text(raw)
    words = len(text.split())
    if words == 0:
        hint = " It may be a scanned PDF." if suffix == ".pdf" else ""
        return Extracted("", 0, f"No readable text found.{hint}")
    if words < SHORT_SUBMISSION_WORDS:
        return Extracted(text, words, f"Very short submission ({words} words).")
    return Extracted(text, words)

def _pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ExtractionError("This PDF is password-protected.")
        return "\n\n".join(page.get_text() for page in doc)

def _docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")

def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()

def student_name_from_filename(filename: str) -> str:
    name = re.sub(r"[_\-.]+", " ", _path(filename).stem)
    name = re.sub(r"\s+", " ", name).strip()
    if name.islower():
        name = name.title()
    return name or "Student"

def assign_student_ids(names: list[str]) -> list[str]:
    taken: set[str] = set()
    ids = []
    for name in names:
        base = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:60] or "student"
        candidate, n = base, 2
        while candidate in taken:
            candidate = f"{base}-{n}"
            n += 1
        taken.add(candidate)
        ids.append(candidate)
    return ids