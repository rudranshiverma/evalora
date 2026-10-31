import io

import docx
import pymupdf
import pytest

from app.ingest import assign_student_ids, clean_text, extract_text, student_name_from_filename

ESSAY = "Renewable energy adoption in India has accelerated because solar costs fell sharply. " * 3


def flat(text: str) -> str:
    return " ".join(text.split())


def make_pdf(text: str | None = None, password: str | None = None) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text)
    if password:
        data = doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password
        )
    else:
        data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    document = docx.Document()
    document.add_paragraph(ESSAY)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Source"
    table.rows[0].cells[1].text = "IEA report"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_pdf() -> None:
    result = extract_text("essay.pdf", make_pdf(ESSAY))
    assert result.warning is None
    assert result.word_count == len(ESSAY.split())
    assert "solar costs fell sharply" in flat(result.text)


def test_docx_includes_tables() -> None:
    result = extract_text("essay.DOCX", make_docx())
    assert result.warning is None
    assert "solar costs fell sharply" in result.text
    assert "Source | IEA report" in result.text


@pytest.mark.parametrize(
    ("data", "expected"),
    [
        ("\ufeffCafé notes".encode(), "Café notes"),
        ("Café notes".encode("cp1252"), "Café notes"),
    ],
)
def test_txt_encodings(data: bytes, expected: str) -> None:
    assert extract_text("notes.txt", data).text == expected


@pytest.mark.parametrize(
    ("filename", "data", "warning"),
    [
        ("scan.pdf", make_pdf(), "scanned PDF"),
        ("broken.pdf", b"not really a pdf", "corrupted"),
        ("locked.pdf", make_pdf(ESSAY, password="secret"), "password-protected"),
        ("photo.png", b"\x89PNG", "Unsupported file type .png"),
        ("README", b"text", "with no extension"),
        ("short.txt", b"Too short to grade.", "Very short submission (4 words)"),
    ],
)
def test_warnings(filename: str, data: bytes, warning: str) -> None:
    assert warning in (extract_text(filename, data).warning or "")


def test_clean_text() -> None:
    assert clean_text("a  \r\nb\r\n\r\n\r\n\r\nc\x00") == "a\nb\n\nc"


@pytest.mark.parametrize(
    ("filename", "name"),
    [
        ("asha_rao.pdf", "Asha Rao"),
        ("Ravi-Kumar Essay.docx", "Ravi Kumar Essay"),
        ("C:\\Users\\teacher\\meera.txt", "Meera"),
        ("___.pdf", "Student"),
    ],
)
def test_student_names(filename: str, name: str) -> None:
    assert student_name_from_filename(filename) == name


def test_student_ids_are_unique_and_safe() -> None:
    ids = assign_student_ids(["Asha Rao", "asha rao", "आशा", "Student"])
    assert ids == ["asha-rao", "asha-rao-2", "student", "student-2"]
